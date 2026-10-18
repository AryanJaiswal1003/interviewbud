import io
import re
import zipfile

from docx import Document
from pypdf import PdfReader

from .. import llm
from .schemas import CandidateProfile


MAX_PDF_PAGES = 20
MAX_DOCX_UNZIPPED = 40 * 1024 * 1024  # a resume is a few hundred KB; this stops zip bombs


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(f"That PDF has {len(reader.pages)} pages; resumes over {MAX_PDF_PAGES} pages aren't supported")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif name.endswith(".docx"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNZIPPED:
                    raise ValueError("That Word file expands to an unreasonable size")
        except zipfile.BadZipFile:
            raise ValueError("That isn't a valid Word (.docx) file")
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError("Unsupported resume format; upload PDF, DOCX, TXT or MD")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 200:
        raise ValueError("Could not read enough text from the resume (is it a scanned image?)")
    return text[:30000]


SYSTEM = """You extract a structured candidate profile from a resume for a technical screening interview.
Rules:
- Use only what the resume states. Never invent employers, dates, numbers or technologies. Use "" or [] when absent.
- claimed_impact and claims_to_verify must quote or closely paraphrase the resume's own specific claims
  (numbers, scale, ownership, novel technology). 4-10 claims, prioritising the most impressive/checkable ones.
- seniority_estimate reflects demonstrated scope, not just years.
- conversation_hooks: 2-5 genuine job-relevant details for rapport. Never include age, nationality, religion,
  family, health or other protected characteristics.
- Exclude protected-category information entirely (photos, birth dates, marital status, nationality) even if present."""


async def build_profile(resume_text: str) -> CandidateProfile:
    return await llm.structured(CandidateProfile, SYSTEM, llm.untrusted("resume", resume_text), temperature=0)
